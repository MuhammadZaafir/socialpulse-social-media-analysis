"""
reports.py
----------
Everything that creates downloadable files:
  * CSV template and Excel template (with instructions)
  * cleaned data as CSV / Excel
  * summary text, HTML report and PDF report

All reports are built from the data and filters passed in - nothing is hard-coded.
Nothing is sent to any external service; files are created in memory.
"""

import io
from datetime import date
from html import escape
from xml.sax.saxutils import escape as xml_escape

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.shapes import Drawing
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from modules import metrics
from modules.data_validation import COLUMN_GUIDE, STANDARD_COLUMNS
from modules.metrics import format_number, format_percent

REPORT_TITLE = "SocialPulse - Social Media Engagement Analysis Report"
INPUT_STEPS = [
    "Download the template.",
    "Open it in Excel.",
    "Enter data from your own authorised social media analytics.",
    "Use dates in YYYY-MM-DD format where possible.",
    "Enter numeric metrics as numbers, not text.",
    "Leave unavailable metrics blank instead of inventing values.",
    "Save as CSV UTF-8 or Excel (.xlsx).",
    "Upload the file to SocialPulse.",
]


# ----------------------------------------------------------------------------
# Templates and data downloads
# ----------------------------------------------------------------------------
def template_csv_bytes():
    """CSV template: the 12 headings plus 3 blank rows. Blank rows are ignored by the app."""
    lines = [",".join(STANDARD_COLUMNS)] + [",".join([""] * (len(STANDARD_COLUMNS) - 1))] * 3
    return ("\n".join(lines) + "\n").encode("utf-8-sig")   # utf-8-sig so Excel opens it correctly


def template_excel_bytes():
    """Excel template with a 'Template' sheet (to fill in) and an 'Instructions' sheet."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Template"
    sheet.append(STANDARD_COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
    for column_letter in "ABCDEFGHIJKL":
        sheet.column_dimensions[column_letter].width = 16
    sheet.column_dimensions["E"].width = 40
    for row in range(2, 5):                       # three blank rows; date column formatted as a date
        sheet.cell(row=row, column=2).number_format = "yyyy-mm-dd"
    sheet.freeze_panes = "A2"

    guide = workbook.create_sheet("Instructions")
    guide.append(["How to use this template"])
    guide["A1"].font = Font(bold=True, size=14)
    for number, step in enumerate(INPUT_STEPS, start=1):
        guide.append([f"{number}. {step}"])
    guide.append([])
    guide.append(["Column", "Required?", "Meaning", "Example"])
    header_row = guide.max_row
    for cell in guide[header_row]:
        cell.font = Font(bold=True)
    for name, (required, meaning, example) in COLUMN_GUIDE.items():
        guide.append([name, required, meaning, example])
    guide.append([])
    guide.append(["Note: example values above are only illustrations. Do not leave example data in the Template sheet."])
    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["B"].width = 12
    guide.column_dimensions["C"].width = 70
    guide.column_dimensions["D"].width = 28
    for row in guide.iter_rows(min_row=1, max_row=guide.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=False, vertical="top")
    return _workbook_bytes(workbook)


def _workbook_bytes(workbook):
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _export_copy(df):
    """Copy of a DataFrame prepared for export (dates as YYYY-MM-DD, missing values stay blank)."""
    out = df.copy()
    if "post_date" in out.columns:
        out["post_date"] = pd.to_datetime(out["post_date"]).dt.strftime("%Y-%m-%d")
    return out


def dataframe_to_csv_bytes(df):
    return _export_copy(df).to_csv(index=False).encode("utf-8-sig")


def dataframe_to_excel_bytes(df, sheet_name="Data", extra_sheets=None):
    """Excel file with the data and, optionally, extra sheets {name: DataFrame}."""
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        _export_copy(df).to_excel(writer, sheet_name=sheet_name, index=False)
        for name, table in (extra_sheets or {}).items():
            _export_copy(table).to_excel(writer, sheet_name=name[:31], index=False)
    return buffer.getvalue()


# ----------------------------------------------------------------------------
# Report content (one dictionary used by the text, HTML and PDF reports)
# ----------------------------------------------------------------------------
def build_report_data(df, source_label, period_text, filters_text, rate_col, rate_note,
                      insight_result, data_warnings, total_posts_in_file):
    """Collect everything the report needs from the CURRENT filtered data."""
    summary = metrics.summarize(df, rate_col)
    pair = metrics.best_and_worst_posts(df, rate_col) if not df.empty else None

    rows = [("Posts analysed", format_number(summary["posts"])), ("Posts in complete dataset", format_number(total_posts_in_file))]
    for key, label in (("likes", "Total likes"), ("comments", "Total comments"), ("shares", "Total shares"),
                       ("impressions", "Total impressions"), ("reach", "Total reach (sum of post-level reach)"),
                       ("clicks", "Total clicks")):
        value = summary["total_" + key]
        rows.append((label, format_number(value) if value is not None else metrics.NOT_AVAILABLE))
    rate_label = metrics.RATE_LABELS.get(rate_col, "Engagement rate")
    rows.append((f"Average {rate_label}",
                 format_percent(summary["avg_rate"]) + f" (mean of {summary['posts_with_rate']} posts)"
                 if summary["avg_rate"] is not None else metrics.NOT_AVAILABLE))

    group_tables = {}
    for column, title in (("platform", "By platform"), ("content_type", "By content type"), ("topic", "By topic")):
        table = metrics.group_summary(df, column, rate_col)
        if not table.empty and len(table) > 1:
            group_tables[title] = (column, table)

    top_posts = pd.DataFrame()
    basis_label = ""
    if pair is not None:
        ranked = metrics.rank_posts(df, pair["basis"]).head(5)
        keep = [c for c in ("rank", "post_id", "post_date", "platform", "content_type", pair["basis"]) if c in ranked.columns]
        top_posts = ranked[keep]
        basis_label = pair["basis_label"]

    return {
        "title": REPORT_TITLE,
        "generated": date.today().isoformat(),
        "source_label": source_label,
        "period": period_text,
        "filters": filters_text,
        "posts": len(df),
        "available_metrics": [c for c in ("likes", "comments", "shares", "impressions", "reach", "clicks") if c in df.columns],
        "rate_label": rate_label,
        "methodology": rate_note,
        "summary_rows": rows,
        "group_tables": group_tables,
        "top_posts": top_posts,
        "top_basis": basis_label,
        "findings": insight_result["findings"],
        "recommendations": insight_result["recommendations"],
        "limitations": insight_result["limitations"],
        "warnings": list(data_warnings),
    }


def build_summary_text(report):
    """Plain-text summary of metrics, findings and recommendations."""
    lines = [report["title"], "=" * len(report["title"]), f"Generated: {report['generated']}",
             f"Data source: {report['source_label']}", f"Period: {report['period']}", f"Filters: {report['filters']}", "",
             "KEY NUMBERS"]
    lines += [f"- {label}: {value}" for label, value in report["summary_rows"]]
    lines += ["", "METHOD", report["methodology"], "", "FINDINGS"]
    lines += [f"- {f['title']} (based on {f['posts']} posts): {f['text']}" for f in report["findings"]] or ["- None"]
    lines += ["", "RECOMMENDATIONS"]
    lines += [f"- {r['text']} [Basis: {r['basis']}]" for r in report["recommendations"]]
    lines += ["", "DATA-QUALITY WARNINGS"] + ([f"- {w}" for w in report["warnings"]] or ["- None"])
    lines += ["", "LIMITATIONS"] + [f"- {t}" for t in report["limitations"]]
    return "\n".join(lines).encode("utf-8")


# ----------------------------------------------------------------------------
# HTML report
# ----------------------------------------------------------------------------
def _html_table(df):
    shown = df.copy()
    if "post_date" in shown.columns:
        shown["post_date"] = pd.to_datetime(shown["post_date"]).dt.strftime("%Y-%m-%d")
    return shown.to_html(index=False, escape=True, na_rep="Not available", float_format=lambda v: f"{v:,.2f}", border=0)


def build_html_report(report):
    """A self-contained HTML report (open in any browser; use Print > Save as PDF if you want a PDF)."""
    e = escape
    parts = [
        "<!DOCTYPE html><html lang='en'><head><meta charset='utf-8'>",
        f"<title>{e(report['title'])}</title>",
        "<style>body{font-family:Arial,sans-serif;max-width:900px;margin:30px auto;color:#222;line-height:1.5}"
        "h1{color:#1F4E79}h2{border-bottom:2px solid #1F4E79;padding-bottom:4px;margin-top:30px}"
        "table{border-collapse:collapse;width:100%;font-size:14px}th{background:#1F4E79;color:#fff;text-align:left}"
        "td,th{padding:6px 8px;border-bottom:1px solid #ddd}.note{color:#555;font-size:13px}</style></head><body>",
        f"<h1>{e(report['title'])}</h1>",
        f"<p class='note'>Generated: {e(report['generated'])}<br>Data source: {e(report['source_label'])}<br>"
        f"Analysis period: {e(report['period'])}<br>Filters: {e(report['filters'])}<br>"
        f"Posts analysed: {report['posts']}<br>Available metrics: {e(', '.join(report['available_metrics']))}</p>",
        "<h2>Key numbers</h2><table>",
    ]
    parts += [f"<tr><td>{e(label)}</td><td>{e(value)}</td></tr>" for label, value in report["summary_rows"]]
    parts += ["</table>", "<h2>Engagement calculation methodology</h2>", f"<p>{e(report['methodology'])}</p>"]
    for title, (column, table) in report["group_tables"].items():
        keep = [c for c in (column, "posts", "avg_engagement_rate", "avg_interactions") if c in table.columns]
        parts += [f"<h2>Performance {e(title.lower())}</h2>", _html_table(table[keep])]
    if not report["top_posts"].empty:
        parts += [f"<h2>Top posts (ranked by {e(report['top_basis'])})</h2>", _html_table(report["top_posts"])]
    parts.append("<h2>Main findings</h2><ul>")
    parts += [f"<li><b>{e(f['title'])}</b> (based on {f['posts']} posts): {e(f['text'])}</li>" for f in report["findings"]] or ["<li>None</li>"]
    parts.append("</ul><h2>Recommendations</h2><ul>")
    parts += [f"<li>{e(r['text'])} <span class='note'>Basis: {e(r['basis'])}</span></li>" for r in report["recommendations"]]
    parts.append("</ul><h2>Data-quality warnings</h2><ul>")
    parts += [f"<li>{e(w)}</li>" for w in report["warnings"]] or ["<li>None</li>"]
    parts.append("</ul><h2>Dataset limitations</h2><ul>")
    parts += [f"<li>{e(t)}</li>" for t in report["limitations"]]
    parts.append("</ul></body></html>")
    return "\n".join(parts).encode("utf-8")


# ----------------------------------------------------------------------------
# PDF report (reportlab)
# ----------------------------------------------------------------------------
def _pdf_text(text):
    """Escape text for reportlab and replace characters the built-in PDF font cannot draw."""
    safe = str(text).encode("cp1252", "replace").decode("cp1252")
    return xml_escape(safe)


def _pdf_table(rows, col_widths=None, header=True):
    table = Table(rows, colWidths=col_widths, repeatRows=1 if header else 0)
    style = [("FONTSIZE", (0, 0), (-1, -1), 8.5), ("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey),
             ("VALIGN", (0, 0), (-1, -1), "TOP"), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F6FA")])]
    if header:
        style += [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white)]
    table.setStyle(TableStyle(style))
    return table


def _bar_chart(labels, values):
    """Small bar chart drawn with reportlab (no image libraries needed)."""
    drawing = Drawing(440, 190)
    chart = VerticalBarChart()
    chart.x, chart.y, chart.width, chart.height = 45, 45, 380, 120
    chart.data = [[0 if pd.isna(v) else float(v) for v in values]]
    chart.categoryAxis.categoryNames = [_pdf_text(str(label))[:18] for label in labels]
    chart.categoryAxis.labels.angle = 25
    chart.categoryAxis.labels.boxAnchor = "ne"
    chart.categoryAxis.labels.fontSize = 7
    chart.valueAxis.labels.fontSize = 7
    chart.valueAxis.valueMin = 0
    chart.bars[0].fillColor = colors.HexColor("#1F4E79")
    drawing.add(chart)
    return drawing


def build_pdf_report(report):
    """Create the PDF report in memory and return its bytes."""
    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["BodyText"], fontSize=9.5, leading=13)
    small = ParagraphStyle("small", parent=body, fontSize=8.5, textColor=colors.HexColor("#555555"))
    h1 = ParagraphStyle("h1", parent=styles["Title"], fontSize=18, textColor=colors.HexColor("#1F4E79"))
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=12, textColor=colors.HexColor("#1F4E79"), spaceBefore=10)
    P = lambda text, style=body: Paragraph(_pdf_text(text), style)

    story = [P(report["title"], h1),
             P(f"Generated: {report['generated']}   |   Posts analysed: {report['posts']}", small),
             P(f"Data source: {report['source_label']}", small),
             P(f"Analysis period: {report['period']}", small),
             P(f"Filters: {report['filters']}", small),
             P(f"Available metrics: {', '.join(report['available_metrics']) or 'none'}", small),
             P("Key numbers", h2)]
    story.append(_pdf_table([["Measure", "Value"]] + [[P(a, body), P(b, body)] for a, b in report["summary_rows"]],
                            col_widths=[7.5 * cm, 9 * cm]))
    story += [P("Engagement calculation methodology", h2), P(report["methodology"])]

    for title, (column, table) in report["group_tables"].items():
        value_col = "avg_engagement_rate" if "avg_engagement_rate" in table.columns else "avg_interactions"
        value_name = report["rate_label"] if value_col == "avg_engagement_rate" else "Avg interactions per post"
        rows = [[column.replace("_", " ").title(), "Posts", value_name]]
        for _, r in table.iterrows():
            shown = format_percent(r[value_col]) if value_col == "avg_engagement_rate" else format_number(r[value_col], 1)
            rows.append([P(r[column]), int(r["posts"]), shown])
        story += [P(f"Performance {title.lower()}", h2), _pdf_table(rows, col_widths=[6 * cm, 2.5 * cm, 8 * cm])]
        if len(table) <= 8:
            story.append(_bar_chart(table[column], table[value_col]))

    if not report["top_posts"].empty:
        top = report["top_posts"]
        rows = [[c.replace("_", " ").title() for c in top.columns]]
        for _, r in top.iterrows():
            rows.append([P(pd.Timestamp(v).strftime("%Y-%m-%d") if c == "post_date" else
                           (f"{v:,.2f}" if isinstance(v, float) else v)) for c, v in r.items()])
        story += [P(f"Top posts (ranked by {report['top_basis']})", h2), _pdf_table(rows)]

    story.append(P("Main findings", h2))
    for f in report["findings"] or [{"title": "None", "text": "No findings available.", "posts": 0}]:
        story.append(P(f"- {f['title']} (based on {f['posts']} posts): {f['text']}"))
    story.append(P("Recommendations", h2))
    for r in report["recommendations"]:
        story.append(P(f"- {r['text']}  (Basis: {r['basis']})"))
    story.append(P("Data-quality warnings", h2))
    for w in report["warnings"] or ["None"]:
        story.append(P(f"- {w}"))
    story.append(P("Dataset limitations", h2))
    for t in report["limitations"]:
        story.append(P(f"- {t}"))
    story += [Spacer(1, 10), P("Generated locally by SocialPulse from the data you provided. No data was sent to external services.", small)]

    buffer = io.BytesIO()
    SimpleDocTemplate(buffer, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.8 * cm,
                      bottomMargin=1.8 * cm, title=report["title"]).build(story)
    return buffer.getvalue()
